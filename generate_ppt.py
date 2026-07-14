from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Paths to images
    title_img = "/Users/akhil/.gemini/antigravity-ide/brain/f162e204-0dfc-4e8e-9f05-9bff5aeabef2/qaforge_title_bg_1781791385335.png"
    ai_models_img = "/Users/akhil/.gemini/antigravity-ide/brain/f162e204-0dfc-4e8e-9f05-9bff5aeabef2/qaforge_ai_models_1781791399260.png"
    dashboard_img = "/Users/akhil/.gemini/antigravity-ide/brain/f162e204-0dfc-4e8e-9f05-9bff5aeabef2/qaforge_dashboard_1781791412727.png"

    # Slide dimensions
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    # --- Slide 1: Title Slide ---
    slide_layout = prs.slide_layouts[6] # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Add background image
    try:
        slide.shapes.add_picture(title_img, 0, 0, width=slide_width, height=slide_height)
    except Exception as e:
        print("Could not load title image:", e)
        
    # Add a dark rectangle to make text readable
    shape = slide.shapes.add_shape(
        1, # MSO_SHAPE.RECTANGLE
        0, slide_height - Inches(3.5), slide_width, Inches(3.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(20, 20, 20)
    shape.line.fill.background()
    
    # Add Text over shape
    txBox = slide.shapes.add_textbox(Inches(0.5), slide_height - Inches(3), slide_width - Inches(1), Inches(2.5))
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = "QAForge"
    p.font.bold = True
    p.font.size = Pt(64)
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = "AI-Powered Testing Suite (Gemini Edition)"
    p2.font.size = Pt(32)
    p2.font.color.rgb = RGBColor(100, 200, 255)

    p3 = tf.add_paragraph()
    p3.text = "Automated Test Generation, Execution, and Reporting"
    p3.font.size = Pt(24)
    p3.font.color.rgb = RGBColor(200, 200, 200)

    # --- Slide 2: What is QAForge? ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "What is QAForge?"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "A revolutionary testing tool powered by Google Gemini AI."
    
    bullet_points = [
        "Generate comprehensive test cases effortlessly using AI.",
        "Execute test suites with real-time pass/fail tracking.",
        "Export rich test reports containing KPIs and beautiful charts."
    ]
    for pt in bullet_points:
        p = tf.add_paragraph()
        p.text = pt
        p.level = 1

    # --- Slide 3: Supported AI Models ---
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), slide_width - Inches(1), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "The Power of Gemini AI Models"
    p.font.size = Pt(40)
    p.font.bold = True
    
    # Add Image on the right
    try:
        # Scale image to fit height nicely
        slide.shapes.add_picture(ai_models_img, Inches(5), Inches(1.5), width=Inches(4.5))
    except Exception as e:
        print("Could not load AI models image:", e)
        
    # Add content on the left
    txBox_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(5))
    tf_c = txBox_content.text_frame
    tf_c.word_wrap = True
    
    p = tf_c.paragraphs[0]
    p.text = "QAForge integrates four specialized Gemini models:"
    p.font.size = Pt(24)
    p.font.bold = True
    
    models = [
        "gemini-2.0-flash: Fast URL & text generation",
        "gemini-1.5-pro: Deep analysis & complex suites",
        "gemini-1.5-flash: Image/screenshot vision",
        "gemini-2.0-flash-exp: Video recording analysis"
    ]
    for m in models:
        p = tf_c.add_paragraph()
        p.text = m
        p.font.size = Pt(20)
        p.level = 0
        p.space_before = Pt(14)

    # --- Slide 4: Key Features ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Key Features"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Capabilities designed for modern QA teams:"
    
    features = [
        "Generate Tests from Anywhere: URLs, Videos, Images, or Text",
        "Cross-Platform Support: Web, Mobile (iOS/Android), Desktop",
        "Live Execution Engine: Real-time feedback and tracking",
        "Flexible Model Selection: Choose the perfect AI for the job",
        "Zero-Friction Config: Easy setup directly in the UI",
        "Extensible: Support for Playwright, Selenium, and Appium"
    ]
    for f in features:
        p = tf.add_paragraph()
        p.text = f
        p.level = 1

    # --- Slide 5: Comprehensive Reporting ---
    slide_layout = prs.slide_layouts[6] # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), slide_width - Inches(1), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Reporting & Analytics Dashboard"
    p.font.size = Pt(40)
    p.font.bold = True

    # Image in the center
    try:
        # Scale to max width with some margin
        slide.shapes.add_picture(dashboard_img, Inches(1), Inches(1.5), width=slide_width - Inches(2))
    except Exception as e:
        print("Could not load dashboard image:", e)

    # Bottom text box
    shape = slide.shapes.add_shape(1, 0, slide_height - Inches(1.5), slide_width, Inches(1.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(20, 20, 20)
    shape.line.fill.background()
    
    txBox2 = slide.shapes.add_textbox(Inches(0.5), slide_height - Inches(1.2), slide_width - Inches(1), Inches(1))
    tf2 = txBox2.text_frame
    p = tf2.paragraphs[0]
    p.text = "Export professional HTML, Excel (with Charts), and JSON reports."
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.size = Pt(24)
    p.alignment = PP_ALIGN.CENTER

    # --- Slide 6: Conclusion ---
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Ready to Elevate Your QA?"
    subtitle = slide.placeholders[1]
    subtitle.text = "QAForge: Intelligence Meets Automation\n\nQuestions & Demo"

    prs.save("QAForge_Presentation.pptx")
    print("Attractive presentation generated at QAForge_Presentation.pptx")

if __name__ == "__main__":
    create_presentation()
